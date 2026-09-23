"""Kayıt listesi / düzenleme ve puanlama view testleri."""
import io
import json
import os

import pytest
from django.contrib.auth.models import User
from django.urls import reverse
from openpyxl import load_workbook

from optik import store
from optik.services import core, scoring, snapshot

# 1x1 PNG
PNG_BYTES = bytes.fromhex(
    "89504e470d0a1a0a0000000d4948445200000001000000010806000000"
    "1f15c4890000000d49444154789c6360000002000100e221bc330000000049454e44ae426082"
)


# ── Kurulum ─────────────────────────────────────────────────────────────

def _rec(record_id, student_no, booklet, answers, **extra):
    d = {"record_id": record_id, "student_no": student_no, "booklet": booklet,
         "class_code": "01", "answers": answers, "seating": None}
    d.update(extra)
    return d


@pytest.fixture
def other(db):
    return User.objects.create_user(username="baskasi", password="x")


@pytest.fixture
def test_id(user):
    tid = core.create_test(course_code="MAT101", course_name="Matematik", exam_type="Vize",
                           booklets=["A", "B"], num_items=5, created_by=user)
    core.set_answer_key(tid, "A", {"Q1": "A", "Q2": "B", "Q3": "C", "Q4": "D", "Q5": "E"})
    core.set_answer_key(tid, "B", {"Q1": "E", "Q2": "D", "Q3": "C", "Q4": "B", "Q5": "A"})
    return tid


@pytest.fixture
def batch_id(test_id, user):
    bid = core.create_batch(test_id, "Şube 1", user)
    core.save_record(bid, "r1", _rec("r1", "1001", "A", ["A", "B", "C", "D", "E"]))
    core.save_record(bid, "r2", _rec("r2", "1002", "B", ["E", "D", "", "!", "A"]))
    core.save_record(bid, "r3", _rec("r3", "10?3", "A", ["A", "A", "A", "A", "A"]))
    return bid


@pytest.fixture
def owner_client(client, user):
    client.force_login(user)
    return client


@pytest.fixture
def other_client(client, other):
    client.force_login(other)
    return client


def _png(batch_id, name="form.png"):
    path = os.path.join(store.get_upload_dir(batch_id), name)
    with open(path, "wb") as f:
        f.write(PNG_BYTES)
    return path


# ── Yetki: her view için sahip / başkası / anonim ───────────────────────

def _get_urls(test_id, batch_id):
    return [
        reverse("optik:records", args=[test_id]),
        reverse("optik:batch_detail", args=[batch_id]),
        reverse("optik:record_detail", args=[batch_id, "r1"]),
        reverse("optik:record_edit", args=[batch_id, "r1"]),
        reverse("optik:scoring", args=[test_id]),
        reverse("optik:open_scoring", args=[batch_id]),
        reverse("optik:batch_score", args=[batch_id]),
        reverse("optik:batch_scores", args=[batch_id]),
        reverse("optik:student_score", args=[batch_id, "1001"]),
        reverse("optik:export_scores", args=[batch_id]),
        reverse("optik:approve_form", args=[batch_id]),
    ]


def _post_urls(batch_id):
    return [
        reverse("optik:records_delete_selected", args=[batch_id]),
        reverse("optik:record_delete", args=[batch_id, "r1"]),
        reverse("optik:record_save", args=[batch_id, "r1"]),
        reverse("optik:do_scoring", args=[batch_id]),
        reverse("optik:save_open_scores", args=[batch_id]),
        reverse("optik:approve", args=[batch_id]),
    ]


@pytest.mark.django_db
def test_owner_can_open_all_pages(owner_client, test_id, batch_id, user):
    scoring.score_batch(batch_id, user)
    for url in _get_urls(test_id, batch_id):
        resp = owner_client.get(url)
        assert resp.status_code in (200, 302), url


@pytest.mark.django_db
def test_other_user_forbidden(other_client, test_id, batch_id):
    urls = _get_urls(test_id, batch_id) + [
        reverse("optik:record_image", args=[batch_id, "r1"]),
        reverse("optik:record_overlay", args=[batch_id, "r1"]),
    ]
    for url in urls:
        assert other_client.get(url).status_code == 403, url
    for url in _post_urls(batch_id):
        assert other_client.post(url).status_code == 403, url
    assert store.get_record(batch_id, "r1")


@pytest.mark.django_db
def test_anonymous_redirected_to_login(client, test_id, batch_id):
    for url in _get_urls(test_id, batch_id) + _post_urls(batch_id):
        resp = client.get(url) if url in _get_urls(test_id, batch_id) else client.post(url)
        assert resp.status_code == 302, url
        assert reverse("login") in resp["Location"]


@pytest.mark.django_db
def test_admin_can_access(client, test_id, batch_id):
    admin = User.objects.create_user(username="yonetici", password="x", is_staff=True)
    client.force_login(admin)
    assert client.get(reverse("optik:records", args=[test_id])).status_code == 200


@pytest.mark.django_db
def test_unknown_ids_404(owner_client, batch_id):
    assert owner_client.get(reverse("optik:records", args=["test_yok"])).status_code == 404
    assert owner_client.get(reverse("optik:record_detail", args=[batch_id, "yok"])).status_code == 404
    assert owner_client.get(reverse("optik:snapshot_detail", args=[batch_id, "v9"])).status_code == 404


# ── Kayıt listesi ───────────────────────────────────────────────────────

@pytest.mark.django_db
def test_records_page_sorting_and_warnings(owner_client, test_id, batch_id):
    core.save_record(batch_id, "r4", _rec("r4", "1001", "A", ["A", "B", "C", "D", "E"]))
    resp = owner_client.get(reverse("optik:records", args=[test_id]))
    assert resp.status_code == 200
    rows = resp.context["rows"]
    # hatalı (10?3) → mükerrer (1001 ×2) → boş/çift (1002)
    assert [r["record_id"] for r in rows][0] == "r3"
    assert {rows[1]["record_id"], rows[2]["record_id"]} == {"r1", "r4"}
    assert rows[3]["record_id"] == "r2"
    texts = [w["text"] for w in rows[3]["warnings"]]
    assert "1 Boş" in texts and "1 Çift" in texts
    assert resp.context["dup_count"] == 1
    assert resp.context["warning_count"] == 3
    assert resp.context["temp_warning_count"] == 1
    assert resp.context["ok_count"] == 0
    html = resp.content.decode()
    assert "Mükerrer Kayıt Tespit Edildi!" in html
    assert "Öğr. No Hatalı" in html
    assert resp.context["wizard_step"] == 4


@pytest.mark.django_db
def test_records_page_without_batch(owner_client, user):
    tid = core.create_test("X", "Y", "Final", ["A"], 3, user)
    resp = owner_client.get(reverse("optik:records", args=[tid]))
    assert resp.status_code == 200
    assert "Henüz form yüklenmemiş" in resp.content.decode()


@pytest.mark.django_db
def test_batch_detail_redirect(owner_client, test_id, batch_id):
    resp = owner_client.get(reverse("optik:batch_detail", args=[batch_id]))
    assert resp.status_code == 302
    assert resp["Location"] == reverse("optik:records", args=[test_id])


@pytest.mark.django_db
def test_record_delete_json(owner_client, batch_id):
    url = reverse("optik:record_delete", args=[batch_id, "r1"])
    assert owner_client.get(url).status_code == 405
    resp = owner_client.post(url)
    assert resp.json() == {"success": True}
    assert "r1" not in {r["record_id"] for r in store.list_records(batch_id)}
    resp = owner_client.post(url)
    assert resp.status_code == 404 and resp.json()["success"] is False


@pytest.mark.django_db
def test_delete_selected(owner_client, batch_id):
    url = reverse("optik:records_delete_selected", args=[batch_id])
    resp = owner_client.post(url, data=json.dumps({"record_ids": ["r1", "r2", "yok"]}),
                             content_type="application/json")
    assert resp.json() == {"success": True, "deleted": 2}
    assert store.count_records(batch_id) == 1
    assert core.get_batch(batch_id)["record_count"] == 1

    resp = owner_client.post(url, data=json.dumps({"record_ids": []}), content_type="application/json")
    assert resp.status_code == 400 and resp.json()["error"] == "Kayıt seçilmedi"
    resp = owner_client.post(url, data="bozuk", content_type="application/json")
    assert resp.status_code == 400


# ── Kayıt detay / düzenleme / kaydetme ──────────────────────────────────

@pytest.mark.django_db
def test_record_detail_grid_and_images(owner_client, test_id, batch_id):
    rec = store.get_record(batch_id, "r2")
    rec["image_path"] = _png(batch_id)
    core.save_record(batch_id, "r2", rec)

    resp = owner_client.get(reverse("optik:record_detail", args=[batch_id, "r2"]))
    assert resp.status_code == 200
    cells = resp.context["cells"]
    assert [c["css"] for c in cells] == ["optik-cell-correct", "optik-cell-correct",
                                         "optik-cell-blank", "optik-cell-multi", "optik-cell-correct"]
    assert resp.context["total_answered"] == 4 and resp.context["total_blank"] == 1
    html = resp.content.decode()
    assert reverse("optik:record_image", args=[batch_id, "r2"]) in html
    assert "base64" not in html
    assert resp.context["has_image"] and not resp.context["has_overlay"]


@pytest.mark.django_db
def test_record_detail_return_url_validated(owner_client, batch_id):
    url = reverse("optik:record_detail", args=[batch_id, "r1"])
    resp = owner_client.get(url, {"return_url": "https://evil.example.com/x"})
    assert resp.context["return_url"] == ""
    resp = owner_client.get(url, {"return_url": "/optik/"})
    assert resp.context["back_url"] == "/optik/"
    assert "return_url=%2Foptik%2F" in resp.context["edit_url"]


@pytest.mark.django_db
def test_record_edit_page(owner_client, batch_id):
    resp = owner_client.get(reverse("optik:record_edit", args=[batch_id, "r2"]))
    assert resp.status_code == 200
    rows = resp.context["rows"]
    assert len(rows) == 5
    assert rows[0]["status"] == "ok"
    assert [o["label"] for o in rows[0]["options"]] == ["A", "B", "C", "D", "E", "-"]
    assert rows[2]["options"][-1]["selected"] is True  # boş cevap
    html = resp.content.decode()
    assert 'name="ans_1"' in html and 'name="student_no"' in html
    assert "csrfmiddlewaretoken" in html


@pytest.mark.django_db
def test_record_save_writes_overrides_and_history(owner_client, user, batch_id):
    url = reverse("optik:record_save", args=[batch_id, "r3"])
    resp = owner_client.post(url, {
        "student_no": " 1003 ", "class_code": "02", "booklet": "B",
        "ans_1": "e", "ans_2": "d", "ans_3": "", "ans_4": "B", "ans_5": "A", "ans_x": "Z",
        "return_url": "",
    })
    assert resp.status_code == 302
    assert resp["Location"] == reverse("optik:record_detail", args=[batch_id, "r3"])
    rec = store.get_record(batch_id, "r3")
    assert rec["student_no"] == "1003" and rec["class_code"] == "02" and rec["booklet"] == "B"
    assert rec["answers"] == ["E", "D", "", "B", "A"]
    assert rec["overrides"] == {"student_no": "1003", "class_code": "02", "booklet": "B"}
    hist = rec["override_history"][-1]
    assert hist["user_id"] == user.pk and hist["student_no"] == "1003" and hist["booklet"] == "B"
    assert "timestamp" in hist and "updated_at" in rec
    # student_no sütunu da güncellendi
    from optik.models import OptikRecord
    assert OptikRecord.objects.get(record_id="r3").student_no == "1003"


@pytest.mark.django_db
def test_record_save_keeps_dict_answers_and_return_url(owner_client, batch_id):
    core.save_record(batch_id, "r5", _rec("r5", "2000", "A", {"Q1": "A", "Q2": "B"}))
    url = reverse("optik:record_save", args=[batch_id, "r5"])
    resp = owner_client.post(url, {"student_no": "2000", "class_code": "", "booklet": "A",
                                   "ans_1": "C", "ans_2": "", "return_url": "/optik/test/x/records/"})
    assert resp["Location"] == "/optik/test/x/records/"
    assert store.get_record(batch_id, "r5")["answers"] == {"Q1": "C", "Q2": ""}
    resp = owner_client.post(url, {"student_no": "2000", "booklet": "A",
                                   "return_url": "https://evil.example.com/"})
    assert resp["Location"] == reverse("optik:record_detail", args=[batch_id, "r5"])


# ── Dosya sunumu ────────────────────────────────────────────────────────

@pytest.mark.django_db
def test_record_image_and_overlay_serving(owner_client, batch_id):
    img = _png(batch_id, "orijinal.png")
    ov = _png(batch_id, "overlay.png")
    rec = store.get_record(batch_id, "r1")
    rec["image_path"] = img
    core.save_record(batch_id, "r1", rec)

    resp = owner_client.get(reverse("optik:record_image", args=[batch_id, "r1"]))
    assert resp.status_code == 200 and resp["Content-Type"] == "image/png"
    assert b"".join(resp.streaming_content) == PNG_BYTES
    # overlay yok → orijinal görüntüye düşer
    resp = owner_client.get(reverse("optik:record_overlay", args=[batch_id, "r1"]))
    assert resp.status_code == 200
    resp.close()

    rec["overlay_path"] = ov
    core.save_record(batch_id, "r1", rec)
    resp = owner_client.get(reverse("optik:record_overlay", args=[batch_id, "r1"]))
    assert resp.status_code == 200 and resp["Content-Type"] == "image/png"
    resp.close()


@pytest.mark.django_db
def test_file_serving_path_traversal_guard(owner_client, test_id, batch_id, tmp_path):
    outside = tmp_path / "gizli.png"
    outside.write_bytes(PNG_BYTES)
    rec = store.get_record(batch_id, "r1")
    rec["image_path"] = str(outside)
    rec["overlay_path"] = os.path.join(store.get_upload_dir(batch_id), "..", "..", "..", "gizli.png")
    core.save_record(batch_id, "r1", rec)
    assert owner_client.get(reverse("optik:record_image", args=[batch_id, "r1"])).status_code == 404
    assert owner_client.get(reverse("optik:record_overlay", args=[batch_id, "r1"])).status_code == 404

    # Başka batch'in dosyası da sunulmaz
    other_bid = core.create_batch(test_id, "Diğer", None)
    rec["image_path"] = _png(other_bid)
    rec.pop("overlay_path")
    core.save_record(batch_id, "r1", rec)
    assert owner_client.get(reverse("optik:record_image", args=[batch_id, "r1"])).status_code == 404


# ── Puanlama ────────────────────────────────────────────────────────────

@pytest.mark.django_db
def test_scoring_end_to_end(owner_client, test_id, batch_id):
    url = reverse("optik:scoring", args=[test_id])
    resp = owner_client.get(url)
    assert resp.status_code == 200
    assert resp.context["show_form"] and resp.context["records_count"] == 3
    assert resp.context["wizard_step"] == 5
    assert reverse("optik:do_scoring", args=[batch_id]) in resp.content.decode()

    resp = owner_client.post(reverse("optik:do_scoring", args=[batch_id]))
    assert resp.status_code == 302 and resp["Location"] == url
    assert store.load_scores(batch_id)
    assert core.get_batch(batch_id)["status"] == "SCORED"

    resp = owner_client.get(url)
    assert "optik/scoring/results.html" in [t.name for t in resp.templates]
    rows = {r["student_no"]: r for r in resp.context["rows"]}
    assert rows["1001"]["total"] == 100.0 and rows["1001"]["passed"]
    assert [c["css"] for c in rows["1002"]["cells"]] == [
        "optik-cell-correct", "optik-cell-correct", "optik-cell-blank", "optik-cell-wrong",
        "optik-cell-correct"]
    assert rows["1001"]["record_id"] == "r1"
    assert resp.context["item_ids"] == ["Q1", "Q2", "Q3", "Q4", "Q5"]
    html = resp.content.decode()
    assert "Puanlama Sonuçları" in html and "Yeniden Puanla" in html
    for fmt in ("tsv", "csv", "xlsx"):
        assert f"fmt={fmt}&amp;detail=1" in html
    assert "publish" not in html and "Yayınla" not in html

    # ?rescore=1 → form yeniden
    resp = owner_client.get(url, {"rescore": "1"})
    assert resp.context["show_form"]


@pytest.mark.django_db
def test_scoring_missing_key_and_duplicates(owner_client, user, test_id, batch_id):
    core.save_record(batch_id, "r9", _rec("r9", "1001", "C", ["A"]))
    resp = owner_client.get(reverse("optik:scoring", args=[test_id]))
    assert resp.context["missing_keys"] == ["C"]
    assert "Cevap Anahtarı Eksik" in resp.content.decode()

    core.delete_record(batch_id, "r9")
    core.save_record(batch_id, "r9", _rec("r9", "1001", "A", ["A"]))
    resp = owner_client.get(reverse("optik:scoring", args=[test_id]))
    assert resp.context["duplicates"] == ["1001"]
    assert "Mükerrer Kayıt Uyarısı" in resp.content.decode()


@pytest.mark.django_db
def test_do_scoring_error_page(owner_client, test_id, user):
    bid = core.create_batch(test_id, "Boş", user)
    resp = owner_client.post(reverse("optik:do_scoring", args=[bid]))
    assert resp.status_code == 400
    assert "Puanlama Hatası" in resp.content.decode()


@pytest.fixture
def open_test(user):
    tid = core.create_test("FIZ", "Fizik", "Vize", ["A"], 3, user,
                           question_types={"MCQ": 2, "TF": 0, "OPEN": 1},
                           points_by_type={"MCQ": 80.0, "TF": 0.0, "OPEN": 20.0})
    core.set_answer_key(tid, "A", {"Q1": "A", "Q2": "B"})
    bid = core.create_batch(tid, "Şube", user)
    core.save_record(bid, "o1", _rec("o1", "3001", "A", ["A", "B"]))
    core.save_record(bid, "o2", _rec("o2", "3002", "A", ["A", "C"]))
    return tid, bid


@pytest.mark.django_db
def test_scoring_with_open_items(owner_client, open_test):
    tid, bid = open_test
    resp = owner_client.get(reverse("optik:scoring", args=[tid]))
    assert resp.context["open_items"] == [{"item_id": "Q3", "points": 20.0}]
    html = resp.content.decode()
    assert 'name="o1_Q3"' in html and 'max="20.0"' in html

    resp = owner_client.post(reverse("optik:do_scoring", args=[bid]), {"o1_Q3": "15", "o2_Q3": "50"})
    assert resp.status_code == 302
    assert store.get_record(bid, "o1")["open_scores"] == {"Q3": 15.0}
    assert store.get_record(bid, "o2")["open_total"] == 20.0  # max'a kırpıldı
    recs = {r["record_id"]: r for r in scoring.get_batch_scores(bid)["records"]}
    assert recs["o1"]["summary"]["total_points"] == 95.0
    assert recs["o2"]["summary"]["total_points"] == 60.0


@pytest.mark.django_db
def test_open_scoring_page_and_save(owner_client, user, open_test, test_id, batch_id):
    tid, bid = open_test
    scoring.score_batch(bid, user)
    resp = owner_client.get(reverse("optik:open_scoring", args=[bid]))
    assert resp.status_code == 200 and len(resp.context["open_rows"]) == 2

    resp = owner_client.post(reverse("optik:save_open_scores", args=[bid]), {"o1_Q3": "10", "o2_Q3": "abc"})
    assert resp["Location"] == reverse("optik:scoring", args=[tid])
    rec = store.get_record(bid, "o1")
    assert rec["open_scores"] == {"Q3": 10.0} and rec["open_total"] == 10.0
    assert store.get_record(bid, "o2")["open_scores"] == {}
    recs = {r["record_id"]: r for r in scoring.get_batch_scores(bid)["records"]}
    assert recs["o1"]["summary"]["total_points"] == 90.0

    # açık uçlu soru yoksa puanlama sayfasına yönlenir
    resp = owner_client.get(reverse("optik:open_scoring", args=[batch_id]))
    assert resp["Location"] == reverse("optik:scoring", args=[test_id])


@pytest.mark.django_db
def test_batch_score_pages(owner_client, user, test_id, batch_id):
    resp = owner_client.get(reverse("optik:batch_score", args=[batch_id]))
    assert resp.status_code == 200 and resp.context["title"] == "Otomatik Puanlama"

    resp = owner_client.get(reverse("optik:batch_scores", args=[batch_id]))
    assert "Henüz puanlama yapılmamış" in resp.content.decode()

    scoring.score_batch(batch_id, user)
    resp = owner_client.get(reverse("optik:batch_scores", args=[batch_id]))
    totals = [r["total"] for r in resp.context["rows"]]
    assert totals == sorted(totals, reverse=True)
    html = resp.content.decode()
    assert reverse("kopya:panel", args=[batch_id]) not in html

    resp = owner_client.get(reverse("optik:student_score", args=[batch_id, "1002"]))
    assert resp.status_code == 200
    labels = [r["label"] for r in resp.context["item_rows"]]
    assert labels == ["Doğru", "Doğru", "Boş", "Çoklu", "Doğru"]
    assert owner_client.get(reverse("optik:student_score", args=[batch_id, "9999"])).status_code == 404

    empty = core.create_batch(test_id, "Boş", user)
    resp = owner_client.get(reverse("optik:batch_score", args=[empty]))
    assert "Önce form yüklemelisiniz" in resp.content.decode()


# ── Dışa aktarma ────────────────────────────────────────────────────────

@pytest.mark.django_db
@pytest.mark.parametrize("fmt,ctype", [
    ("tsv", "text/tab-separated-values; charset=utf-8"),
    ("csv", "text/csv; charset=utf-8"),
    ("xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
])
@pytest.mark.parametrize("detail", [0, 1])
def test_export_formats(owner_client, user, batch_id, fmt, ctype, detail):
    scoring.score_batch(batch_id, user)
    resp = owner_client.get(reverse("optik:export_scores", args=[batch_id]), {"fmt": fmt, "detail": detail})
    assert resp.status_code == 200
    assert resp["Content-Type"] == ctype
    suffix = "summary" if detail == 0 else "detail"
    assert resp["Content-Disposition"] == f'attachment; filename="{batch_id}_scores_{suffix}.{fmt}"'
    assert len(resp.content) > 0
    if fmt == "xlsx":
        ws = load_workbook(io.BytesIO(resp.content)).active
        header = [c.value for c in ws[1]]
    else:
        sep = "\t" if fmt == "tsv" else ","
        header = resp.content.decode().splitlines()[0].split(sep)
    assert header[0] == "student_no"
    assert ("Q5" in header) == (detail == 1)


@pytest.mark.django_db
def test_export_invalid_format(owner_client, batch_id):
    resp = owner_client.get(reverse("optik:export_scores", args=[batch_id]), {"fmt": "pdf"})
    assert resp.status_code == 400
    resp = owner_client.get(reverse("optik:export_scores", args=[batch_id]), {"detail": "x"})
    assert resp.status_code == 200 and resp.content == b"No data"


# ── Onay + snapshot ─────────────────────────────────────────────────────

@pytest.mark.django_db
def test_approve_creates_snapshot(owner_client, user, batch_id):
    resp = owner_client.get(reverse("optik:approve_form", args=[batch_id]))
    assert resp.status_code == 200 and "Onayla ve Snapshot Al" in resp.content.decode()

    resp = owner_client.post(reverse("optik:approve", args=[batch_id]), {"reason": "İlk onay"})
    assert resp.status_code == 302
    assert resp["Location"] == reverse("optik:batch_detail", args=[batch_id])
    batch = core.get_batch(batch_id)
    assert batch["status"] == "APPROVED" and batch["approved_by"] == user.pk
    meta = snapshot.get_snapshot(batch_id, "v1")["metadata"]
    assert meta["reason"] == "İlk onay"

    resp = owner_client.get(reverse("optik:snapshot_detail", args=[batch_id, "v1"]))
    assert resp.status_code == 200
    html = resp.content.decode()
    assert "Snapshot v1" in html and "İlk onay" in html
    assert "1001" in html and "1002" in html and "10?3" in html
    assert resp.context["record_count"] == 3
    assert resp.context["created_by"] == "Şükrü Işık"


@pytest.mark.django_db
def test_approve_default_reason(owner_client, batch_id):
    owner_client.post(reverse("optik:approve", args=[batch_id]))
    assert snapshot.get_snapshot(batch_id, "v1")["metadata"]["reason"] == "Batch approved"
