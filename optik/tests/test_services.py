"""optik.services iş mantığı testleri."""
import csv
import io
import math
import os

import pytest
from openpyxl import Workbook, load_workbook

from optik import store
from optik.models import OptikTest
from optik.services import (
    background, core, export, omr, override, scoring, snapshot,
)


# ── Yardımcılar ─────────────────────────────────────────────────────────

def _make_test(user, **kw):
    params = dict(course_code="MAT101", course_name="Matematik", exam_type="Vize",
                  booklets=["A", "B"], num_items=5, created_by=user)
    params.update(kw)
    return core.create_test(**params)


def _rec(record_id, student_no, booklet, answers, **extra):
    d = {"record_id": record_id, "student_no": student_no, "booklet": booklet,
         "answers": answers, "seating": None}
    d.update(extra)
    return d


@pytest.fixture
def test_id(user):
    tid = _make_test(user)
    core.set_answer_key(tid, "A", {"Q1": "A", "Q2": "B", "Q3": "*", "Q4": "C", "Q5": "D"})
    core.set_answer_key(tid, "B", {"Q1": "E", "Q2": "D", "Q3": "C", "Q4": "B", "Q5": "A"})
    return tid


@pytest.fixture
def batch_id(test_id, user):
    return core.create_batch(test_id, "Şube 1", user)


# ── Test / cevap anahtarı / batch ───────────────────────────────────────

@pytest.mark.django_db
def test_create_test_and_auto_item_points(user):
    tid = _make_test(user, academic_year="2025-2026", semester="Güz")
    test = core.get_test(tid)
    assert test["created_by"] == user.id
    assert test["academic_year"] == "2025-2026" and test["semester"] == "Güz"
    assert list(test["items"]) == ["Q1", "Q2", "Q3", "Q4", "Q5"]
    assert all(math.isclose(i["points"], 20.0) for i in test["items"].values())
    obj = OptikTest.objects.get(test_id=tid)
    assert obj.owner == user and obj.semester == "Güz"


@pytest.mark.django_db
def test_auto_item_points_mixed_types(user):
    tid = _make_test(user, num_items=5, question_types={"MCQ": 4, "TF": 0, "OPEN": 1},
                     points_by_type={"MCQ": 80.0, "TF": 0.0, "OPEN": 20.0})
    items = core.get_test(tid)["items"]
    assert items["Q1"] == {"type": "MCQ", "points": 20.0}
    assert items["Q5"] == {"type": "OPEN", "points": 20.0}
    assert core.get_test(tid)["points_total"] == 100.0


@pytest.mark.django_db
def test_answer_keys_and_excel_import(user, test_id):
    assert core.get_test(test_id)["answer_keys"]["A"]["Q3"] == "*"

    wb = Workbook()
    ws = wb.active
    ws.append(["item", "booklet_A", "B", "puan"])
    ws.append(["1", "c", "D", 30])
    ws.append(["Q2", "A", "x", 10])
    buf = io.BytesIO()
    wb.save(buf)
    result = core.import_answer_keys_from_excel(test_id, buf.getvalue())
    assert result["success"] and set(result["booklets_updated"]) == {"A", "B"}
    test = core.get_test(test_id)
    assert test["answer_keys"]["A"] == {"Q1": "C", "Q2": "A"}
    assert test["answer_keys"]["B"] == {"Q1": "D"}
    assert test["items"]["Q1"]["points"] == 30.0


@pytest.mark.django_db
def test_create_batch_and_listing(user, test_id, batch_id):
    batch = core.get_batch(batch_id)
    assert batch["test_id"] == test_id and batch["status"] == "CREATED"
    assert batch["created_by"] == user.id and batch["test_name"] == "MAT101 - Vize"
    assert [b["batch_id"] for b in core.list_test_batches(test_id)] == [batch_id]
    assert core.get_test_batch_id(test_id) == batch_id

    core.approve_batch(batch_id, user)
    batch = core.get_batch(batch_id)
    assert batch["status"] == "APPROVED" and batch["approved_by"] == user.id
    assert "published" not in batch


@pytest.mark.django_db
def test_records_save_list_delete(batch_id):
    for i in range(3):
        core.save_record(batch_id, f"r{i}", _rec(f"r{i}", f"10{i}", "A", ["A"] * 5))
    assert core.get_record_count(batch_id) == 3
    assert {r["record_id"] for r in core.get_batch_records(batch_id)} == {"r0", "r1", "r2"}

    assert core.delete_record(batch_id, "r0") is True
    assert core.delete_record(batch_id, "r0") is False
    assert core.get_batch(batch_id)["record_count"] == 2
    assert core.delete_records(batch_id, ["r1", "r2", "yok"]) == 2
    assert core.get_record_count(batch_id) == 0


def _set_override(batch_id, record_id, field, value):
    rec = store.get_record(batch_id, record_id)
    rec.setdefault("overrides", {})[field] = value
    store.save_record(batch_id, record_id, rec)


# ── Puanlama ────────────────────────────────────────────────────────────

@pytest.mark.django_db
def test_score_batch_rules(user, test_id, batch_id):
    core.save_record(batch_id, "r1", _rec("r1", "1001", "A", ["A", "C", "", "!", "D"]))
    core.save_record(batch_id, "r2", _rec("r2", "1002", "B", ["E", "D", "C", "B", "A"]))
    core.save_record(batch_id, "r3", _rec("r3", "1003", "", ["A"] * 5))      # kitapçık yok → atlanır
    core.save_record(batch_id, "r4", _rec("r4", "1004", "C", ["A"] * 5))     # anahtar yok → atlanır

    scores = scoring.score_batch(batch_id, user)
    by_id = {r["record_id"]: r for r in scores["records"]}
    assert set(by_id) == {"r1", "r2"}

    r1 = by_id["r1"]
    st = {k: v["status"] for k, v in r1["item_scores"].items()}
    assert st == {"Q1": "correct", "Q2": "wrong", "Q3": "correct", "Q4": "multi", "Q5": "correct"}
    assert r1["summary"]["correct"] == 3 and r1["summary"]["wrong"] == 2 and r1["summary"]["blank"] == 0
    assert r1["summary"]["total_points"] == 60.0 and r1["summary"]["percentage"] == 60.0
    assert by_id["r2"]["summary"]["total_points"] == 100.0

    assert scores["scored_by"] == user.id
    assert scores["metadata"]["booklets"] == ["A", "B"]
    assert scores["metadata"]["total_records"] == 2
    assert scores["class_statistics"]["mean"] == 80.0
    assert core.get_batch(batch_id)["status"] == "SCORED"
    assert scoring.get_batch_scores(batch_id)["records"] == scores["records"]
    assert scoring.get_student_score(batch_id, "1002")["record_id"] == "r2"


@pytest.mark.django_db
def test_score_single_record_blank_and_key_blank():
    res = scoring.score_single_record(
        {"record_id": "x", "answers": {"Q1": "?", "Q2": "B", "Q3": ""}},
        {"Q1": "A", "Q2": "", "Q3": "C"}, {"Q1": 10, "Q2": 10, "Q3": 10})
    st = {k: v["status"] for k, v in res["item_scores"].items()}
    assert st == {"Q1": "blank", "Q2": "not_graded", "Q3": "blank"}
    assert res["summary"]["blank"] == 2 and res["summary"]["max_points"] == 20


@pytest.mark.django_db
def test_score_batch_all_skipped_no_name_error(user, test_id, batch_id):
    core.save_record(batch_id, "r1", _rec("r1", "1", "", ["A"] * 5))
    scores = scoring.score_batch(batch_id, user)
    assert scores["records"] == [] and scores["class_statistics"] == {}
    assert scores["metadata"]["booklet"] is None and scores["metadata"]["total_records"] == 0


@pytest.mark.django_db
def test_score_batch_empty_raises(user, batch_id):
    with pytest.raises(ValueError):
        scoring.score_batch(batch_id, user)


@pytest.mark.django_db
def test_score_batch_uses_overrides(user, test_id, batch_id):
    core.save_record(batch_id, "r1", _rec("r1", "1001", "!", ["E", "D", "C", "B", "A"]))
    _set_override(batch_id, "r1", "booklet", "B")
    eff = override.get_effective_result(batch_id, "r1")
    assert eff["effective"]["booklet"] == "B" and eff["has_overrides"]
    scores = scoring.score_batch(batch_id, user)
    assert scores["records"][0]["summary"]["total_points"] == 100.0
    with pytest.raises(store.NotFound):
        override.get_effective_result(batch_id, "yok")


@pytest.mark.django_db
def test_update_scores_json_open_scores_duplicate_student_no(user):
    tid = _make_test(user, booklets=["A"], num_items=3,
                     question_types={"MCQ": 2, "TF": 0, "OPEN": 1},
                     points_by_type={"MCQ": 80.0, "TF": 0.0, "OPEN": 20.0})
    core.set_answer_key(tid, "A", {"Q1": "A", "Q2": "B"})
    bid = core.create_batch(tid, "b", user)
    core.save_record(bid, "r1", _rec("r1", "555", "A", ["A", "B"]))
    core.save_record(bid, "r2", _rec("r2", "555", "A", ["A", "C"]))
    scoring.score_batch(bid, user)

    for rid, pts in (("r1", 5), ("r2", 15)):
        rec = store.get_record(bid, rid)
        rec["open_scores"] = {"Q3": pts}
        rec["open_total"] = pts
        store.save_record(bid, rid, rec)
    scoring.update_scores_json(bid)

    by_id = {r["record_id"]: r for r in scoring.get_batch_scores(bid)["records"]}
    assert by_id["r1"]["open_total"] == 5 and by_id["r2"]["open_total"] == 15
    assert by_id["r1"]["summary"]["total_points"] == 85.0     # 80 + 5
    assert by_id["r2"]["summary"]["total_points"] == 55.0     # 40 + 15
    assert by_id["r1"]["summary"]["max_points"] == 100.0
    stats = scoring.get_batch_scores(bid)["class_statistics"]
    assert stats["mean"] == 70.0
    assert "updated_at" in scoring.get_batch_scores(bid)


# ── Snapshot ────────────────────────────────────────────────────────────

@pytest.mark.django_db
def test_snapshot_includes_records(user, batch_id):
    core.save_record(batch_id, "r1", _rec("r1", "1001", "A", ["A"] * 5, image_path="/x/1.png"))
    core.save_record(batch_id, "r2", _rec("r2", "1002", "B", ["B"] * 5))
    core.save_record(batch_id, "r3", {"record_id": "r3", "status": "UPLOADED"})
    _set_override(batch_id, "r2", "student_no", "2002")

    v1 = snapshot.create_snapshot(batch_id, user, "onay")
    assert v1 == "v1"
    snap = snapshot.get_snapshot(batch_id, "v1")
    assert snap["metadata"]["record_count"] == 2
    assert snap["metadata"]["created_by"] == user.id
    by_id = {r["record_id"]: r for r in snap["records"]}
    assert set(by_id) == {"r1", "r2"}
    assert by_id["r2"]["effective"]["student_no"] == "2002"
    assert by_id["r2"]["override_fields"] == ["student_no"]

    with pytest.raises(store.NotFound):
        snapshot.get_snapshot(batch_id, "v9")
    assert snapshot.create_snapshot(batch_id, user) == "v2"


# ── Export ──────────────────────────────────────────────────────────────

@pytest.fixture
def wide_batch(user):
    """30 maddelik puanlanmış batch (detay export'ta 35 sütun)."""
    tid = _make_test(user, booklets=["A"], num_items=30)
    core.set_answer_key(tid, "A", {f"Q{i}": "A" for i in range(1, 31)})
    bid = core.create_batch(tid, "geniş", user)
    core.save_record(bid, "r1", _rec("r1", "1001", "A", ["A"] * 30))
    core.save_record(bid, "r2", _rec("r2", "1002", "A", ["B"] * 15 + ["A"] * 15))
    scoring.score_batch(bid, user)
    return bid


@pytest.mark.django_db
def test_export_tsv_csv(wide_batch):
    tsv = export.export_to_tsv(wide_batch, 0).decode("utf-8").split("\n")
    assert tsv[0] == "student_no\ttotal_score\tclass_code\tseating_cell_id\tbooklet"
    assert tsv[1].startswith("1001\t100.0")

    rows = list(csv.reader(io.StringIO(export.export_to_csv(wide_batch, 1).decode("utf-8"))))
    assert len(rows[0]) == 35 and rows[0][-1] == "Q30"
    assert rows[2][1] == "50.0"


@pytest.mark.django_db
def test_export_csv_quoting(wide_batch):
    scores = store.load_scores(wide_batch)
    scores["records"][0]["student_no"] = 'a,b"c'
    store.save_scores(wide_batch, scores)
    rows = list(csv.reader(io.StringIO(export.export_to_csv(wide_batch, 0).decode("utf-8"))))
    assert rows[1][0] == 'a,b"c' and len(rows[1]) == 5


@pytest.mark.django_db
def test_export_xlsx_more_than_26_columns(wide_batch):
    data = export.export_to_xlsx(wide_batch, 1)
    ws = load_workbook(io.BytesIO(data)).active
    header = [c.value for c in ws[1]]
    assert len(header) == 35 and header[-1] == "Q30"
    assert ws.column_dimensions["AI"].width == 8        # 35. sütun (Q30)
    assert ws.column_dimensions["A"].width == 15
    assert ws.cell(row=2, column=2).value == 100.0
    assert export.get_export_filename(wide_batch, "xlsx", 1).endswith("_scores_detail.xlsx")


@pytest.mark.django_db
def test_export_no_scores(batch_id):
    assert export.export_to_tsv(batch_id) == b"No data"
    assert export.export_to_csv(batch_id) == b"No data"
    assert load_workbook(io.BytesIO(export.export_to_xlsx(batch_id))).active["A1"].value == "No data"


# ── OMR ve arka plan işleme ─────────────────────────────────────────────

STUDENT_ANSWERS = list("ABCDEEDCBA") + ["", "C", "", "A", "B", "D", "E", "", "A", "C"]
KEY_ANSWERS = list("ABCDEEDCBAABCDEEDCBA")


@pytest.fixture
def omr_test(user):
    tid = _make_test(user, booklets=["A", "B"], num_items=20)
    bid = core.create_batch(tid, "optik", user)
    return tid, bid


@pytest.mark.django_db
def test_process_form_image_reads_synthetic_form(make_form_image, omr_test):
    tid, bid = omr_test
    path = make_form_image("202412345678", "B", STUDENT_ANSWERS)
    ok, rec, err = omr.process_form_image(path, tid, ["A", "B"], 20,
                                          overlay_dir=store.get_upload_dir(bid))
    assert ok, err
    assert rec["student_no"] == "202412345678"
    assert rec["booklet"] == "B"
    assert rec["answers"] == STUDENT_ANSWERS
    assert rec["warnings"] is None and rec["omr_confidence"] >= 0.9
    assert store.is_inside_batch_dir(bid, rec["overlay_path"]) and os.path.exists(rec["overlay_path"])


@pytest.mark.django_db
def test_process_forms_background_sync(make_form_image, omr_test, tmp_path):
    tid, bid = omr_test
    key = make_form_image("111111111111", "A", KEY_ANSWERS, name="anahtar.png")
    student = make_form_image("202412345678", "B", STUDENT_ANSWERS, name="ogrenci.png")
    broken = tmp_path / "bozuk.png"
    broken.write_bytes(b"resim degil")

    background.process_forms_background(bid, [key, student, str(broken)])

    # Cevap anahtarı formu kayıt değil, anahtar olarak işlenir
    test = core.get_test(tid)
    assert test["answer_keys"]["A"] == {f"Q{i + 1}": a for i, a in enumerate(KEY_ANSWERS)}

    records = core.get_batch_records(bid)
    assert len(records) == 1
    rec = records[0]
    assert rec["student_no"] == "202412345678"
    assert rec["booklet"] == "B"
    assert rec["answers"] == STUDENT_ANSWERS
    for key_name in ("image_path", "overlay_path"):
        assert store.is_inside_batch_dir(bid, rec[key_name]), rec[key_name]
        assert os.path.exists(rec[key_name])
    assert os.path.dirname(rec["image_path"]) == store.get_upload_dir(bid)

    prog = background.get_batch_progress(bid)
    assert prog["status"] == "COMPLETED_WITH_ERRORS" and prog["is_complete"]
    assert (prog["processed"], prog["success"], prog["failed"]) == (3, 2, 1)
    assert prog["progress_percentage"] == 100 and prog["started_at"] and prog["completed_at"]
    assert prog["errors"][0]["file"] == "bozuk.png"
    assert core.get_batch(bid)["record_count"] == 1


@pytest.mark.django_db
def test_process_forms_background_bad_batch_marks_failed(omr_test):
    tid, bid = omr_test
    background.process_forms_background(bid, [], test_id="test_yok")
    prog = background.get_batch_progress(bid)
    assert prog["status"] == "FAILED" and prog["error"]


@pytest.mark.django_db(transaction=True)
def test_start_processing_runs_in_thread(make_form_image, omr_test, user):
    tid, bid = omr_test
    student = make_form_image("202400000001", "A", STUDENT_ANSWERS)
    future = background.start_processing(bid, [student], user)
    future.result(timeout=120)
    prog = background.get_batch_progress(bid)
    assert prog["status"] == "COMPLETED", prog
    assert core.get_batch_records(bid)[0]["student_no"] == "202400000001"
    assert core.get_batch(bid)["processing_started_by"] == user.id
