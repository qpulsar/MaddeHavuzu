"""
Optik çekirdek servis: test, cevap anahtarı, batch ve kayıt işlemleri.

Kaynak: omr_analysis/service.py

Farklar:
- JSON dosyaları yerine `optik.store` (Django modelleri) kullanılır.
- Kullanıcı parametreleri Django `User` nesnesi (veya doğrudan id) alır;
  dict'lerde kullanıcı id'si (int) saklanır.
- Sınav/dönem (exams/schedule/admin) bağlantıları yok: `academic_year` ve
  `semester` düz alan olarak tutulur.
"""
import uuid
from typing import Any, Dict, List, Optional

from django.contrib.auth import get_user_model

from optik import store
from optik.store import NotFound, now_iso  # noqa: F401  (dışarıya da sunulur)

__all__ = ['create_test', 'get_test', 'set_answer_key',
           'set_item_points', 'auto_set_item_points', 'save_test',
           'set_booklet_mapping', 'update_test', 'set_test_approval',
           'create_batch', 'get_batch', 'save_batch', 'list_test_batches',
           'get_test_batch_id', 'update_batch_status', 'approve_batch',
           'get_batch_records', 'get_record_count', 'save_record',
           'delete_record', 'delete_records',
           'import_answer_keys_from_excel',
           'user_id', 'user_display_name']


# ============================================================================
# KULLANICI YARDIMCILARI
# ============================================================================

def user_id(user: Any) -> Optional[int]:
    """User nesnesi / id / None → id (dict'lerde saklanan değer)."""
    if user is None or user == "":
        return None
    if hasattr(user, "pk"):
        return user.pk
    if isinstance(user, dict):  # kaynak modülün user dict'i
        return user.get("user_id") or user.get("id")
    try:
        return int(user)
    except (TypeError, ValueError):
        return None


def user_display_name(user: Any) -> str:
    """User nesnesi veya id için görünen ad (ad soyad, yoksa kullanıcı adı)."""
    if user is None or user == "":
        return ""
    if not hasattr(user, "get_full_name"):
        uid = user_id(user)
        if uid is None:
            return str(user)
        user = get_user_model().objects.filter(pk=uid).first()
        if user is None:
            return str(uid)
    return user.get_full_name() or user.get_username()


# ============================================================================
# TEST İŞLEMLERİ
# ============================================================================

def load_test(test_id: str) -> Dict[str, Any]:
    return store.load_test(test_id)


def save_test(test_id: str, data: Dict[str, Any]) -> None:
    store.save_test(test_id, data)


def create_test(
    course_code: str,
    course_name: str,
    exam_type: str,
    booklets: List[str],
    num_items: int,
    created_by: Any,
    question_types: Dict[str, int] = None,
    points_by_type: Dict[str, float] = None,
    academic_year: str = None,
    semester: str = None,
) -> str:
    """
    Yeni test oluştur.

    Args:
        created_by: Oluşturan kullanıcı (User nesnesi veya id)
        question_types: {"MCQ": 20, "TF": 0, "OPEN": 0}
        points_by_type: {"MCQ": 100, "TF": 0, "OPEN": 0}

    Returns:
        Test ID
    """
    test_id = f"test_{uuid.uuid4().hex[:8]}"

    # Varsayılan: Tüm sorular MCQ, toplam 100 puan
    if question_types is None:
        question_types = {"MCQ": num_items, "TF": 0, "OPEN": 0}

    if points_by_type is None:
        points_by_type = {"MCQ": 100.0, "TF": 0.0, "OPEN": 0.0}

    test_data = {
        "test_id": test_id,
        "course_code": course_code,
        "course_name": course_name,
        "exam_type": exam_type,
        "booklets": booklets,
        "num_items": num_items,
        "created_by": user_id(created_by),
        "created_at": now_iso(),
        "question_counts": question_types,
        "points_by_type_total": points_by_type,
        "points_total": sum(points_by_type.values()),
        "items": {},
        "answer_keys": {},
    }
    if academic_year:
        test_data["academic_year"] = academic_year
    if semester:
        test_data["semester"] = semester

    save_test(test_id, test_data)

    # Madde puanlarını otomatik hesapla ve kaydet
    auto_set_item_points(test_id)

    return test_id


def auto_set_item_points(test_id: str) -> None:
    """
    Test spec'ten otomatik madde puanlarını hesapla ve kaydet.

    - Her soru tipinin toplam puanı var
    - Her maddenin puanı = tip toplam puanı / o tipteki soru sayısı
    """
    test = load_test(test_id)

    question_counts = test.get("question_counts", {"MCQ": test.get("num_items", 50), "TF": 0, "OPEN": 0})
    points_by_type = test.get("points_by_type_total", {"MCQ": 100.0, "TF": 0.0, "OPEN": 0.0})

    def point_per_item(qtype: str) -> float:
        count = question_counts.get(qtype, 0)
        total = points_by_type.get(qtype, 0.0)
        if count <= 0:
            return 0.0
        return total / count

    mcq_point = point_per_item("MCQ")
    tf_point = point_per_item("TF")
    open_point = point_per_item("OPEN")

    items = {}
    item_no = 1
    for qtype, count, point in (
        ("MCQ", question_counts.get("MCQ", 0), mcq_point),
        ("TF", question_counts.get("TF", 0), tf_point),
        ("OPEN", question_counts.get("OPEN", 0), open_point),
    ):
        for _ in range(count):
            items[f"Q{item_no}"] = {"type": qtype, "points": point}
            item_no += 1

    test["items"] = items
    save_test(test_id, test)


def get_test(test_id: str) -> Dict[str, Any]:
    """Test bilgilerini getir (yoksa NotFound)."""
    return load_test(test_id)


def set_answer_key(test_id: str, booklet: str, answers: Dict[str, str]) -> None:
    """Cevap anahtarını kaydet"""
    test = load_test(test_id)
    test.setdefault("answer_keys", {})
    test["answer_keys"][booklet] = answers
    save_test(test_id, test)


def set_item_points(test_id: str, points: Dict[str, float]) -> None:
    """Madde puanlarını kaydet"""
    test = load_test(test_id)
    test.setdefault("items", {})
    for item_id, point in points.items():
        test["items"].setdefault(item_id, {})
        test["items"][item_id]["points"] = point
    save_test(test_id, test)


def set_booklet_mapping(test_id: str, mapping: Dict[str, Dict[str, int]]) -> None:
    """
    Kitapçık madde eşleştirmesini kaydet.

    mapping: {booklet: {item_id: mapped_item_number}}, örn: {"B": {"Q1": 5, ...}}
    """
    test = load_test(test_id)
    test["booklet_mapping"] = mapping
    save_test(test_id, test)


def update_test(test_id: str, updates: Dict) -> None:
    """Test bilgilerini güncelle"""
    test = load_test(test_id)
    for key, value in updates.items():
        test[key] = value
    test["updated_at"] = now_iso()
    save_test(test_id, test)


def set_test_approval(test_id: str, is_approved: bool) -> None:
    """Test onay durumunu ayarla"""
    test = load_test(test_id)
    test["is_approved"] = is_approved
    test["approval_date"] = now_iso() if is_approved else None
    save_test(test_id, test)


# ============================================================================
# BATCH İŞLEMLERİ
# ============================================================================

def create_batch(test_id: str, batch_name: str, created_by: Any) -> str:
    """
    Yeni batch oluştur.

    Args:
        created_by: User nesnesi veya id

    Returns:
        Batch ID
    """
    batch_id = f"batch_{uuid.uuid4().hex[:8]}"

    test = load_test(test_id)

    batch_data = {
        "batch_id": batch_id,
        "test_id": test_id,
        "batch_name": batch_name,
        "test_name": f"{test.get('course_code', '')} - {test.get('exam_type', '')}",
        "created_by": user_id(created_by),
        "created_at": now_iso(),
        "status": "CREATED",
        "record_count": 0
    }

    store.save_batch(batch_id, batch_data)
    return batch_id


def get_batch(batch_id: str) -> Dict[str, Any]:
    """Batch bilgilerini getir (yoksa NotFound)."""
    return store.load_batch(batch_id)


def save_batch(batch_id: str, batch_data: Dict[str, Any]) -> None:
    """Batch bilgilerini kaydet"""
    store.save_batch(batch_id, batch_data)


def list_test_batches(test_id: str) -> List[Dict[str, Any]]:
    """Bir test'in tüm batch'lerini listele (yeniden eskiye)."""
    batches = []
    for bid in store.list_test_batch_ids(test_id):
        try:
            batches.append(store.load_batch(bid))
        except Exception:
            pass
    batches.sort(key=lambda x: x.get("created_at", ""), reverse=True)
    return batches


def get_test_batch_id(test_id: str) -> Optional[str]:
    """Test'in mevcut batch_id'sini döndür. Yoksa None."""
    batches = list_test_batches(test_id)
    if batches:
        return batches[0]["batch_id"]
    return None


def update_batch_status(batch_id: str, status: str) -> None:
    """Batch durumunu güncelle"""
    store.update_batch_fields(batch_id, status=status, updated_at=now_iso())


def approve_batch(batch_id: str, approved_by: Any) -> None:
    """Batch'i onayla (approved_by: User veya id)."""
    store.update_batch_fields(batch_id, status="APPROVED",
                              approved_by=user_id(approved_by),
                              approved_at=now_iso())


# ============================================================================
# RECORD İŞLEMLERİ
# ============================================================================

def get_batch_records(batch_id: str) -> List[Dict[str, Any]]:
    """Batch'deki tüm kayıtları getir"""
    return store.list_records(batch_id)


def get_record_count(batch_id: str) -> int:
    """Batch'deki kayıt sayısı"""
    return store.count_records(batch_id)


def save_record(batch_id: str, record_id: str, record_data: Dict[str, Any]) -> None:
    """Kayıt kaydet"""
    store.save_record(batch_id, record_id, record_data)


def delete_record(batch_id: str, record_id: str) -> bool:
    """Tek kayıt sil. Bulunamazsa False."""
    if not store.delete_record(batch_id, record_id):
        return False
    store.update_batch_fields(batch_id, record_count=get_record_count(batch_id),
                              updated_at=now_iso())
    return True


def delete_records(batch_id: str, record_ids: List[str]) -> int:
    """Birden fazla kaydı sil; silinen kayıt sayısını döndür."""
    deleted = sum(1 for rid in record_ids if store.delete_record(batch_id, rid))
    if deleted > 0:
        store.update_batch_fields(batch_id, record_count=get_record_count(batch_id),
                                  updated_at=now_iso())
    return deleted


# ============================================================================
# ÖZET BİLGİLER
# ============================================================================


# ============================================================================
# EXCEL IMPORT
# ============================================================================

def import_answer_keys_from_excel(test_id: str, file_content: bytes) -> Dict[str, Any]:
    """
    Excel'den cevap anahtarını import et.

    Excel Formatı:
    | item | booklet_A | booklet_B | points |
    | Q1   | A         | C         | 5.0    |

    Returns:
        {"success": True, "booklets_updated": [...], "items_count": 20, "points_updated": True}
    """
    from io import BytesIO
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(file_content))
    ws = wb.active

    test = load_test(test_id)
    booklets = test.get("booklets", ["A"])

    headers = [str(cell.value).strip().lower() if cell.value else "" for cell in ws[1]]

    item_col = None
    points_col = None
    booklet_cols = {}

    for i, h in enumerate(headers):
        if h in ["item", "soru", "madde", "q"]:
            item_col = i
        elif h in ["points", "puan", "point"]:
            points_col = i
        else:
            # booklet_A, booklet_B veya sadece A, B
            for bk in booklets:
                if h == f"booklet_{bk.lower()}" or h == bk.lower():
                    booklet_cols[bk] = i

    if item_col is None:
        raise ValueError("'item' veya 'soru' sütunu bulunamadı")

    answer_keys = {bk: {} for bk in booklet_cols.keys()}
    item_points = {}

    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row[item_col]:
            continue

        item_id = str(row[item_col]).strip()
        if not item_id.startswith("Q"):
            item_id = f"Q{item_id}"

        for bk, col in booklet_cols.items():
            if col < len(row) and row[col]:
                answer = str(row[col]).strip().upper()
                if answer in ["A", "B", "C", "D", "E", "T", "F", "1", "0"]:
                    answer_keys[bk][item_id] = answer

        if points_col is not None and points_col < len(row) and row[points_col]:
            try:
                item_points[item_id] = float(row[points_col])
            except (ValueError, TypeError):
                pass

    test.setdefault("answer_keys", {})

    booklets_updated = []
    for bk, answers in answer_keys.items():
        if answers:
            test["answer_keys"][bk] = answers
            booklets_updated.append(bk)

    points_updated = False
    if item_points:
        test.setdefault("items", {})
        for item_id, points in item_points.items():
            test["items"].setdefault(item_id, {})
            test["items"][item_id]["points"] = points
        points_updated = True

    save_test(test_id, test)

    return {
        "success": True,
        "booklets_updated": booklets_updated,
        "items_count": len(list(answer_keys.values())[0]) if answer_keys else 0,
        "points_updated": points_updated
    }
