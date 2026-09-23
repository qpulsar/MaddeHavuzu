"""
Excel Yükleme Şablonu
=====================
Kaynak: cheating/validation/make_demo_xlsx.py (`make_template`,
`_write_help_sheet`). Simülasyon bağımlılığı ve sys.path düzenlemesi olmadan
yalnızca boş şablon üretimi taşındı.
"""

from __future__ import annotations

import io
import os
from typing import BinaryIO, Union

from openpyxl import Workbook


def _write_help_sheet(ws, n_booklets: int) -> None:
    """Açıklama sayfası."""
    ws.append(["Kopya Analizi — Excel Yukleme Şablon Açıklaması"])
    ws.append([])
    ws.append(["Bu şablon üç sayfadan oluşur."])
    ws.append([])
    ws.append(["1. Yanitlar (zorunlu)"])
    ws.append(["   Öğrenci başı bir satır."])
    ws.append(["   Zorunlu: ogrenci_no ve S1, S2, ... (madde sütunları)"])
    ws.append(["   Opsiyonel: ad_soyad, salon, oturma_kodu, gozetmen_onayi, kitapcik"])
    ws.append([])
    ws.append(["   Yanit formati: A, B, C, D, E veya 1, 2, 3, 4, 5"])
    ws.append(["     - Büyük/küçük harf duyarsız"])
    ws.append(["     - Boş hücre: cevaplanmamış"])
    ws.append(["     - Çoklu işaret: AB, A,B, A|B, ** vb → geçersiz olarak kodlanır"])
    ws.append([])
    ws.append(["   Oturma kodu formatı: A1, B2, ..., H9 (bir büyük harf + bir sayı)"])
    ws.append(["   Salon: aynı salondaki öğrenciler için aynı değer"])
    ws.append([])
    ws.append(["   gozetmen_onayi: E/H, 1/0, TRUE/FALSE"])
    ws.append([])
    ws.append(["2. Anahtar (zorunlu)"])
    ws.append(["   madde_no: 1'den başlayan sıralı numara"])
    ws.append(["   dogru_cevap: A-E veya 1-5"])
    ws.append([])
    ws.append(["3. Kitapcik_Eslestirme (kitapçık > 1 ise zorunlu)"])
    ws.append(["   kitapcik: Yanitlar sayfasındaki kitapcik değeriyle aynı"])
    ws.append(["   pozisyon: 1..n arası, kitapçıkta sorunun sırası"])
    ws.append(["   madde_no: Anahtar sayfasındaki kanonik numara"])
    ws.append(["   Her kitapçık için tam permutasyon (her madde bir kez, hepsi)"])
    ws.append([])
    ws.append(["Sınırlamalar:"])
    ws.append(["   - Maksimum 5000 öğrenci (aşılırsa hata verilir)"])
    ws.append(["   - Öğrenci numarası tekrar edemez"])
    ws.append(["   - Anahtardaki madde sayısı = madde sütunu sayısı olmalı"])
    ws.append([])
    ws.append(["Sık hatalar:"])
    ws.append(["   - ASCII dışı karakter (Kiril 'а' vs Latin 'a')"])
    ws.append(["   - Kitapçık = 2 ama Kitapcik_Eslestirme sayfası eksik"])
    ws.append(["   - Anahtar geçersiz seçenek (F, Z, 6 gibi)"])


def make_template(out: Union[str, BinaryIO], n_items: int = 20,
                  n_booklets: int = 1, n_options: int = 5) -> Union[str, BinaryIO]:
    """Boş şablon dosyası (5 örnek satır + Aciklama).

    `out` bir dosya yolu veya yazılabilir ikili akış (BytesIO) olabilir.
    """
    if isinstance(out, str):
        os.makedirs(os.path.dirname(out) or ".", exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Yanitlar"
    header = ["ogrenci_no", "ad_soyad", "salon", "oturma_kodu",
              "gozetmen_onayi", "kitapcik"]
    header += [f"S{i+1}" for i in range(n_items)]
    ws.append(header)
    # 5 örnek satır
    example_answers = ["A", "B", "C", "D", "E"] * (n_items // 5 + 1)
    for i in range(5):
        row = [
            f"20260000{i+1:03d}",
            f"Ornek Ogrenci {i+1}",
            "Salon-A",
            f"{chr(ord('A') + i // 3)}{(i % 3) + 1}",
            "E",
            "K1" if n_booklets > 1 else "1",
        ]
        row += example_answers[i:i + n_items]
        ws.append(row)

    ws2 = wb.create_sheet("Anahtar")
    ws2.append(["madde_no", "dogru_cevap"])
    for i in range(n_items):
        ws2.append([i + 1, "A"])

    if n_booklets > 1:
        ws3 = wb.create_sheet("Kitapcik_Eslestirme")
        ws3.append(["kitapcik", "pozisyon", "madde_no"])
        for b in range(n_booklets):
            for p in range(n_items):
                ws3.append([f"K{b+1}", p + 1, p + 1])

    ws_desc = wb.create_sheet("Aciklama")
    _write_help_sheet(ws_desc, n_booklets)

    wb.save(out)
    return out


def make_template_bytes(n_items: int = 20, n_booklets: int = 1) -> bytes:
    """Şablonu bellekte üretip bayt olarak döndür (indirme view'ı için)."""
    buf = io.BytesIO()
    make_template(buf, n_items=n_items, n_booklets=n_booklets)
    return buf.getvalue()
